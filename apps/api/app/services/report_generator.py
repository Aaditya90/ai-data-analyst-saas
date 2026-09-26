"""
Turns a dashboard's already-computed widget data into a downloadable
report file (PDF or PPTX). This module never touches the database or
object storage and never re-derives a number — every value it renders was
computed by app/services/aggregation.py (the same code Phase 6's
"render this widget" endpoint uses), so a report always shows the same
numbers the dashboard would if opened right now. That keeps this module
independently testable with plain dicts, the same posture as
automl.py/forecasting.py.

Chart widgets are rendered to a PNG with matplotlib (Agg backend — no
display, safe on a server) and embedded as an image in both output
formats, rather than using either library's native chart objects. One
rendering code path for both formats, and no risk of the "PowerPoint
reports the file as corrupt" failure mode that native chart XML can hit.

"AI narrates, code validates" — same split as every AI feature since
Phase 8. Nothing in this file calls Claude; the executive summary text is
computed by ai_narration.py's narrate_report_summary() from the exact KPI
values the caller passns in, and handed here already-written, or None.
"""

import io
from datetime import datetime

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.chart.data import CategoryChartData  # noqa: F401 (kept for callers wanting native charts later)
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as RLImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

MAX_TABLE_ROWS_PDF = 25
MAX_TABLE_ROWS_PPTX = 12
CHART_FIGSIZE = (7.2, 3.6)
BRAND_COLOR = "#2F5FE0"


class ReportGenerationError(ValueError):
    """Raised for input that can't be turned into a report."""


def render_chart_image(chart_data: dict) -> bytes:
    """chart_data: {"chart_type": str, "labels": [...], "values": [...]},
    the exact shape aggregation.compute_chart() returns. Returns PNG bytes."""
    labels = chart_data.get("labels") or []
    values = chart_data.get("values") or []
    chart_type = chart_data.get("chart_type", "bar")

    fig, ax = plt.subplots(figsize=CHART_FIGSIZE, dpi=150)
    if not labels or not values:
        ax.text(0.5, 0.5, "No data available", ha="center", va="center", fontsize=12, color="#888888")
        ax.axis("off")
    else:
        # compute_chart's output is always a categorical (label, value)
        # series — line vs. bar is a rendering choice, not a different
        # data shape, so every chart_type other than "line" draws as bars.
        if chart_type == "line":
            ax.plot(labels, values, color=BRAND_COLOR, marker="o", linewidth=2)
        else:
            ax.bar(labels, values, color=BRAND_COLOR)
        ax.tick_params(axis="x", rotation=45, labelsize=8)
        ax.tick_params(axis="y", labelsize=8)
        for spine in ("top", "right"):
            ax.spines[spine].set_visible(False)
        fig.tight_layout()

    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", bbox_inches="tight")
    plt.close(fig)
    buffer.seek(0)
    return buffer.getvalue()


def _validate_sections(sections: list[dict]) -> None:
    if not sections:
        raise ReportGenerationError("Dashboard has no widgets to include in a report")
    for section in sections:
        if "widget_type" not in section or "title" not in section or "data" not in section:
            raise ReportGenerationError(
                "Each section needs widget_type, title, and data — got: " + repr(section)
            )


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------


def generate_pdf_report(
    dashboard_name: str, sections: list[dict], summary: str | None = None, generated_at: datetime | None = None
) -> bytes:
    _validate_sections(sections)
    generated_at = generated_at or datetime.utcnow()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )
    styles = getSampleStyleSheet()
    heading_style = ParagraphStyle(
        "SectionHeading", parent=styles["Heading2"], spaceBefore=18, spaceAfter=8
    )
    story = []

    story.append(Paragraph(dashboard_name, styles["Title"]))
    story.append(
        Paragraph(
            f"Generated {generated_at.strftime('%B %d, %Y at %H:%M UTC')}",
            styles["Normal"],
        )
    )
    story.append(Spacer(1, 16))
    if summary:
        story.append(Paragraph("Executive Summary", heading_style))
        story.append(Paragraph(summary, styles["Normal"]))
    story.append(PageBreak())

    for section in sections:
        story.append(Paragraph(section["title"], heading_style))
        widget_type = section["widget_type"]
        data = section["data"]

        if widget_type == "text":
            story.append(Paragraph(data.get("content", "") or "(empty)", styles["Normal"]))
        elif widget_type == "kpi":
            story.append(_pdf_kpi_flowable(data, styles))
        elif widget_type == "table":
            story.append(_pdf_table_flowable(data, styles))
            total_rows = data.get("total_rows")
            shown_rows = min(len(data.get("rows") or []), MAX_TABLE_ROWS_PDF)
            if total_rows and total_rows > shown_rows:
                story.append(Spacer(1, 4))
                story.append(
                    Paragraph(f"Showing {shown_rows} of {total_rows} rows", styles["Italic"])
                )
        elif widget_type == "chart":
            png_bytes = render_chart_image(data)
            image = RLImage(io.BytesIO(png_bytes), width=6.0 * inch, height=3.0 * inch)
            story.append(image)
        else:
            story.append(Paragraph(f"(unsupported widget type: {widget_type})", styles["Normal"]))
        story.append(Spacer(1, 12))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def _pdf_kpi_flowable(data: dict, styles) -> Paragraph:
    value = data.get("value")
    display = "N/A" if value is None else (f"{value:,.2f}" if isinstance(value, float) else str(value))
    kpi_style = ParagraphStyle("KPIValue", parent=styles["Title"], fontSize=32, textColor=colors.HexColor(BRAND_COLOR))
    return Paragraph(display, kpi_style)


def _pdf_table_flowable(data: dict, styles) -> Table:
    columns = data.get("columns") or []
    rows = (data.get("rows") or [])[:MAX_TABLE_ROWS_PDF]
    if not columns:
        return Table([["(no data)"]])

    table_data = [columns] + rows
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND_COLOR)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FC")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------

_SLIDE_W_IN = 13.333
_SLIDE_H_IN = 7.5


def generate_pptx_report(
    dashboard_name: str, sections: list[dict], summary: str | None = None, generated_at: datetime | None = None
) -> bytes:
    _validate_sections(sections)
    generated_at = generated_at or datetime.utcnow()

    prs = Presentation()
    prs.slide_width = Inches(_SLIDE_W_IN)
    prs.slide_height = Inches(_SLIDE_H_IN)
    blank_layout = prs.slide_layouts[6]

    _add_title_slide(prs, blank_layout, dashboard_name, generated_at)
    if summary:
        _add_text_slide(prs, blank_layout, "Executive Summary", summary)

    for section in sections:
        widget_type = section["widget_type"]
        data = section["data"]
        if widget_type == "text":
            _add_text_slide(prs, blank_layout, section["title"], data.get("content", "") or "(empty)")
        elif widget_type == "kpi":
            _add_kpi_slide(prs, blank_layout, section["title"], data)
        elif widget_type == "table":
            _add_table_slide(prs, blank_layout, section["title"], data)
        elif widget_type == "chart":
            _add_chart_slide(prs, blank_layout, section["title"], data)
        else:
            _add_text_slide(prs, blank_layout, section["title"], f"(unsupported widget type: {widget_type})")

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _add_title_slide(prs, layout, dashboard_name: str, generated_at: datetime) -> None:
    slide = prs.slides.add_slide(layout)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(_SLIDE_W_IN - 1.6), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.text = dashboard_name
    tf.paragraphs[0].font.size = Pt(40)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor.from_string(BRAND_COLOR.lstrip("#"))

    subtitle_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(_SLIDE_W_IN - 1.6), Inches(0.6))
    subtitle_box.text_frame.text = f"Generated {generated_at.strftime('%B %d, %Y at %H:%M UTC')}"
    subtitle_box.text_frame.paragraphs[0].font.size = Pt(16)
    subtitle_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def _add_slide_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(_SLIDE_W_IN - 1.2), Inches(0.8))
    tf = box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True


def _add_text_slide(prs, layout, title: str, content: str) -> None:
    slide = prs.slides.add_slide(layout)
    _add_slide_title(slide, title)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(_SLIDE_W_IN - 1.2), Inches(_SLIDE_H_IN - 2.0))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = content
    tf.paragraphs[0].font.size = Pt(18)


def _add_kpi_slide(prs, layout, title: str, data: dict) -> None:
    slide = prs.slides.add_slide(layout)
    _add_slide_title(slide, title)
    value = data.get("value")
    display = "N/A" if value is None else (f"{value:,.2f}" if isinstance(value, float) else str(value))
    box = slide.shapes.add_textbox(Inches(0.6), Inches(2.6), Inches(_SLIDE_W_IN - 1.2), Inches(2.0))
    tf = box.text_frame
    tf.text = display
    p = tf.paragraphs[0]
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = RGBColor.from_string(BRAND_COLOR.lstrip("#"))
    p.alignment = PP_ALIGN.CENTER


def _add_table_slide(prs, layout, title: str, data: dict) -> None:
    slide = prs.slides.add_slide(layout)
    _add_slide_title(slide, title)
    columns = data.get("columns") or []
    rows = (data.get("rows") or [])[:MAX_TABLE_ROWS_PPTX]
    if not columns:
        _add_text_slide_body(slide, "(no data)")
        return

    n_rows, n_cols = len(rows) + 1, len(columns)
    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.6), Inches(1.4), Inches(_SLIDE_W_IN - 1.2), Inches(min(5.6, 0.4 * n_rows))
    )
    table = table_shape.table
    for col_idx, col_name in enumerate(columns):
        cell = table.cell(0, col_idx)
        cell.text = str(col_name)
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(12)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(BRAND_COLOR.lstrip("#"))
        cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(value)
            cell.text_frame.paragraphs[0].font.size = Pt(11)

    total_rows = data.get("total_rows")
    if total_rows and total_rows > len(rows):
        note = slide.shapes.add_textbox(Inches(0.6), Inches(_SLIDE_H_IN - 0.6), Inches(6), Inches(0.4))
        note.text_frame.text = f"Showing {len(rows)} of {total_rows} rows"
        note.text_frame.paragraphs[0].font.size = Pt(10)
        note.text_frame.paragraphs[0].font.italic = True


def _add_text_slide_body(slide, content: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(_SLIDE_W_IN - 1.2), Inches(2))
    box.text_frame.text = content


def _add_chart_slide(prs, layout, title: str, data: dict) -> None:
    slide = prs.slides.add_slide(layout)
    _add_slide_title(slide, title)
    png_bytes = render_chart_image(data)
    image_stream = io.BytesIO(png_bytes)
    pic_width = Inches(_SLIDE_W_IN - 1.6)
    slide.shapes.add_picture(image_stream, Inches(0.8), Inches(1.5), width=pic_width)

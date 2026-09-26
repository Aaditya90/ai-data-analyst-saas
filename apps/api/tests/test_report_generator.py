"""
Run with: python tests/test_report_generator.py

Covers app/services/report_generator.py end to end with synthetic
already-computed widget data (the exact shape aggregation.py's
compute_kpi/compute_table/compute_chart return) — no DB, no object
storage, no network, no API key needed. Verifies the generated files are
structurally valid (readable page/slide counts) via pypdf and python-pptx,
not just "some bytes came out".
"""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from pypdf import PdfReader

from app.services.report_generator import (
    ReportGenerationError,
    generate_pdf_report,
    generate_pptx_report,
    render_chart_image,
)

CHART_SECTION = {
    "widget_type": "chart",
    "title": "Revenue by Region",
    "data": {"chart_type": "bar", "labels": ["North", "South", "East", "West"], "values": [120.0, 95.5, 60.0, 40.0]},
}
KPI_SECTION = {
    "widget_type": "kpi",
    "title": "Total Revenue",
    "data": {"value": 315500.5},
}
TABLE_SECTION = {
    "widget_type": "table",
    "title": "Top Orders",
    "data": {
        "columns": ["order_id", "amount"],
        "rows": [["ORD-1", "100"], ["ORD-2", "250"]],
        "total_rows": 2,
    },
}
TEXT_SECTION = {
    "widget_type": "text",
    "title": "Notes",
    "data": {"content": "Q3 numbers are preliminary pending finance sign-off."},
}

ALL_SECTIONS = [CHART_SECTION, KPI_SECTION, TABLE_SECTION, TEXT_SECTION]


def test_render_chart_image_returns_valid_png():
    png_bytes = render_chart_image(CHART_SECTION["data"])
    assert png_bytes[:8] == b"\x89PNG\r\n\x1a\n"  # PNG magic bytes
    assert len(png_bytes) > 500


def test_render_chart_image_handles_empty_data():
    png_bytes = render_chart_image({"chart_type": "bar", "labels": [], "values": []})
    assert png_bytes[:8] == b"\x89PNG\r\n\x1a\n"


def test_render_chart_image_line_type_does_not_crash():
    png_bytes = render_chart_image({"chart_type": "line", "labels": ["a", "b", "c"], "values": [1, 2, 3]})
    assert png_bytes[:8] == b"\x89PNG\r\n\x1a\n"


def test_generate_pdf_report_produces_valid_pdf_with_all_widget_types():
    pdf_bytes = generate_pdf_report("Sales Dashboard", ALL_SECTIONS, summary="Q3 revenue grew steadily.")
    reader = PdfReader(io.BytesIO(pdf_bytes))
    # Cover page is a forced PageBreak; sections flow freely after that and
    # may share a page, so just confirm the doc isn't empty and every
    # section's content actually made it into the file.
    assert len(reader.pages) >= 2

    full_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Sales Dashboard" in full_text
    assert "Q3 revenue grew steadily." in full_text
    assert "Revenue by Region" in full_text
    assert "Total Revenue" in full_text
    assert "Top Orders" in full_text
    assert "Q3 numbers are preliminary" in full_text


def test_generate_pdf_report_without_summary_still_builds():
    pdf_bytes = generate_pdf_report("No Summary Dashboard", [KPI_SECTION], summary=None)
    reader = PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1


def test_generate_pdf_report_raises_on_no_sections():
    try:
        generate_pdf_report("Empty Dashboard", [], summary=None)
        assert False, "expected ReportGenerationError"
    except ReportGenerationError:
        pass


def test_generate_pdf_report_table_truncation_note_appears():
    big_table_section = {
        "widget_type": "table",
        "title": "Big Table",
        "data": {
            "columns": ["id"],
            "rows": [[str(i)] for i in range(50)],  # more than MAX_TABLE_ROWS_PDF
            "total_rows": 50,
        },
    }
    pdf_bytes = generate_pdf_report("Big Table Dashboard", [big_table_section], summary=None)
    reader = PdfReader(io.BytesIO(pdf_bytes))
    full_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "of 50 rows" in full_text


def test_generate_pptx_report_produces_valid_pptx_with_all_widget_types():
    pptx_bytes = generate_pptx_report("Sales Dashboard", ALL_SECTIONS, summary="Q3 revenue grew steadily.")
    prs = Presentation(io.BytesIO(pptx_bytes))
    # Title slide + summary slide + one slide per section.
    slide_count = len(list(prs.slides))
    assert slide_count == 2 + len(ALL_SECTIONS)

    # Widescreen (13.333in x 7.5in), in EMU (914400 EMU per inch).
    assert abs(prs.slide_width - 12192000) < 2000
    assert abs(prs.slide_height - 6858000) < 2000


def test_generate_pptx_report_without_summary_has_no_extra_slide():
    pptx_bytes = generate_pptx_report("No Summary Dashboard", [KPI_SECTION], summary=None)
    prs = Presentation(io.BytesIO(pptx_bytes))
    # Title slide + one KPI slide, no summary slide.
    assert len(list(prs.slides)) == 2


def test_generate_pptx_report_raises_on_no_sections():
    try:
        generate_pptx_report("Empty Dashboard", [], summary=None)
        assert False, "expected ReportGenerationError"
    except ReportGenerationError:
        pass


def test_generate_pptx_report_table_slide_has_correct_dimensions():
    pptx_bytes = generate_pptx_report("Table Dashboard", [TABLE_SECTION], summary=None)
    prs = Presentation(io.BytesIO(pptx_bytes))
    table_slide = list(prs.slides)[-1]
    tables = [shape for shape in table_slide.shapes if shape.has_table]
    assert len(tables) == 1
    table = tables[0].table
    assert len(table.rows) == len(TABLE_SECTION["data"]["rows"]) + 1  # +1 header row
    assert len(table.columns) == len(TABLE_SECTION["data"]["columns"])


def test_generate_reports_handle_unsupported_widget_type_gracefully():
    weird_section = {"widget_type": "map", "title": "Unsupported", "data": {}}
    pdf_bytes = generate_pdf_report("Weird Dashboard", [weird_section], summary=None)
    assert len(PdfReader(io.BytesIO(pdf_bytes)).pages) >= 1

    pptx_bytes = generate_pptx_report("Weird Dashboard", [weird_section], summary=None)
    assert len(list(Presentation(io.BytesIO(pptx_bytes)).slides)) == 2


if __name__ == "__main__":
    tests = [v for k, v in globals().items() if k.startswith("test_") and callable(v)]
    for test in tests:
        test()
        print(f"PASS: {test.__name__}")
    print(f"\nAll {len(tests)} Phase 11 report generator unit tests passed.")
